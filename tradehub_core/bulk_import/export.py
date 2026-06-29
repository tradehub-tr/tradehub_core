"""Satıcı ürünlerini toplu yükleme ŞABLONUYLA aynı formatta dışa aktarma.

Round-trip akışın eksik ucu: satıcı mevcut ürünlerini indirir → düzenler →
toplu yükleme sihirbazından `upsert` modunda tekrar yükler → ürünler güncellenir
(SKU eşleştiği için çoğalmaz). Sütunlar `download_template` ile birebir aynıdır
(`build_template_columns` paylaşımı) — böylece export çıktısı doğrudan re-import
edilebilir.
"""

import csv
import io

import frappe
from frappe import _

from tradehub_core.api.listing import build_seller_listing_filters
from tradehub_core.bulk_import.api import build_template_columns
from tradehub_core.utils.tenant import _get_seller_profile_for_user

# Senkron export tavanı — üstünde "filtre daralt" uyarısı (timeout/bellek koruması).
EXPORT_MAX_ROWS = 5000

# 0/1 saklanan alanlar → şablon "Yes"/"No" değerleri (importer ikisini de kabul eder).
_BOOL_CANONICALS = frozenset(
	{"sell_in_moq_multiples", "track_inventory", "allow_backorders", "is_free_shipping"}
)

# Parent satırda boş bırakılan varyant bloğu kolonları (yalnızca varyant satırlarında dolar).
_VARIANT_CANONICALS = frozenset(
	{
		"parent_sku",
		"variant_sku",
		"variant_axis_1_type",
		"variant_axis_1_value",
		"variant_axis_2_type",
		"variant_axis_2_value",
		"variant_axis_3_type",
		"variant_axis_3_value",
		"variant_price",
		"variant_stock",
	}
)

# Çocuk tablodan / özel kaynaklı kolonlar (Listing scalar alanı değil).
_NON_SCALAR_CANONICALS = _VARIANT_CANONICALS | {"image_2", "image_3"}

# Listing'den çekilecek scalar alanlar (canonical == fieldname; sku → seller_sku).
_LISTING_SCALAR_FIELDS = (
	"name",
	"seller_sku",
	"title",
	"brand",
	"condition",
	"product_type",
	"product_family",
	"attribute_set",
	"base_price",
	"selling_price",
	"currency",
	"discount_percentage",
	"stock_qty",
	"stock_uom",
	"min_order_qty",
	"max_order_qty",
	"low_stock_threshold",
	"sell_in_moq_multiples",
	"track_inventory",
	"allow_backorders",
	"is_free_shipping",
	"shipping_weight",
	"handling_days",
	"ships_from_country",
	"ships_from_city",
	"country_of_origin",
	"barcode",
	"video_url",
	"short_description",
	"description",
	"primary_image",
)


def _cell(value) -> str:
	"""Hücre değerini stringe çevir (None → boş, bool 0/1 çağıran tarafça işlenir)."""
	if value is None:
		return ""
	return str(value)


def _scalar_value(listing: dict, canonical: str):
	"""Parent satır için bir canonical alanın Listing değeri."""
	if canonical == "sku":
		return listing.get("seller_sku")
	if canonical in _BOOL_CANONICALS:
		return "Yes" if listing.get(canonical) else "No"
	return listing.get(canonical)


def _build_parent_value_map(listing: dict, attrs: dict, gallery: list[str]) -> dict:
	"""Parent (varyantsız/master) satırın canonical → değer haritası."""
	vmap = {}
	# Galeri: Image 1 = primary_image (scalar), Image 2/3 = listing_images sırası.
	vmap["image_2"] = gallery[0] if len(gallery) > 0 else ""
	vmap["image_3"] = gallery[1] if len(gallery) > 1 else ""
	# attr:<code>
	for code, val in attrs.items():
		vmap[f"attr:{code}"] = val
	return vmap


def _build_variant_value_map(parent_sku: str, v: dict) -> dict:
	"""Varyant satırının canonical → değer haritası (çekirdek alanlar boş)."""
	return {
		"parent_sku": parent_sku,
		"variant_sku": v.get("variant_sku") or "",
		"variant_axis_1_type": v.get("attribute_type") or "",
		"variant_axis_1_value": v.get("attribute_value") or "",
		"variant_axis_2_type": v.get("attribute_type_2") or "",
		"variant_axis_2_value": v.get("attribute_value_2") or "",
		"variant_axis_3_type": v.get("attribute_type_3") or "",
		"variant_axis_3_value": v.get("attribute_value_3") or "",
		"variant_price": v.get("variant_price") or "",
		"variant_stock": v.get("variant_stock") or "",
	}


def _row_from_map(canonicals: list[str], value_map: dict, listing: dict | None = None) -> list[str]:
	"""Canonical sırasına göre satır hücrelerini üret.

	`listing` verilmişse (parent satır) scalar alanlar Listing'den; verilmemişse
	(varyant satır) yalnızca value_map'teki kolonlar dolar, gerisi boş.
	"""
	row = []
	for c in canonicals:
		if c in value_map:
			row.append(_cell(value_map[c]))
		elif listing is not None and c not in _NON_SCALAR_CANONICALS and not c.startswith("attr:"):
			row.append(_cell(_scalar_value(listing, c)))
		else:
			row.append("")
	return row


@frappe.whitelist()
def export_seller_listings(
	format="xlsx",
	status=None,
	search=None,
	product_category=None,
	source=None,
	title=None,
	listing_code=None,
	price_min=None,
	price_max=None,
	stock_min=None,
	stock_max=None,
	completeness_min=None,
	completeness_max=None,
	moq_min=None,
	moq_max=None,
):
	"""Satıcının (filtreli) ürünlerini şablon formatında XLSX/CSV olarak indir.

	Filtreler `get_seller_listings` ile birebir aynı (build_seller_listing_filters)
	— ekrandaki aktif filtre/arama sonucu export edilir. Çıktı doğrudan toplu
	yükleme `upsert` moduyla re-import edilebilir (SKU kolonu zorunlu taşınır).
	"""
	if format not in ("xlsx", "csv"):
		frappe.throw(_("Geçersiz format"))

	seller_profile = _get_seller_profile_for_user(frappe.session.user)
	if not seller_profile:
		frappe.throw(_("Mağaza profili bulunamadı."))

	filters, or_filters = build_seller_listing_filters(
		seller_profile,
		status=status,
		source=source,
		product_category=product_category,
		title=title,
		listing_code=listing_code,
		search=search,
		price_min=price_min,
		price_max=price_max,
		stock_min=stock_min,
		stock_max=stock_max,
		completeness_min=completeness_min,
		completeness_max=completeness_max,
		moq_min=moq_min,
		moq_max=moq_max,
	)

	# get_list → permission_query_conditions devrede (tenant izolasyonu çift güvence).
	names = frappe.get_list(
		"Listing",
		filters=filters,
		or_filters=or_filters,
		pluck="name",
		order_by="creation desc",
		limit_page_length=EXPORT_MAX_ROWS + 1,
	)
	if len(names) > EXPORT_MAX_ROWS:
		frappe.throw(
			_("Çok fazla ürün (>{0}). Lütfen filtre/arama ile daraltıp tekrar deneyin.").format(
				EXPORT_MAX_ROWS
			)
		)
	if not names:
		frappe.throw(_("Dışa aktarılacak ürün bulunamadı."))

	# Scalar alanlar tek sorguda (N+1 yok). Şablonda olup Listing'de DB kolonu
	# OLMAYAN alanlar (örn. barcode bazı kurulumlarda yok) atlanır → export'ta
	# boş kalır (upsert'te boş alan korunduğundan round-trip güvenli).
	valid_cols = set(frappe.get_meta("Listing").get_valid_columns())
	fetch_fields = [f for f in _LISTING_SCALAR_FIELDS if f in valid_cols]
	listings = frappe.get_all(
		"Listing",
		filters={"name": ["in", names]},
		fields=fetch_fields,
	)
	by_name = {row["name"]: row for row in listings}

	# Çocuk tablolar toplu çekilip parent'a göre gruplanır (N+1 yok).
	attr_map: dict[str, dict] = {}
	for r in frappe.get_all(
		"Listing Attribute Value",
		filters={"parent": ["in", names], "parenttype": "Listing"},
		fields=["parent", "attribute", "attribute_value"],
	):
		attr_map.setdefault(r["parent"], {})[r["attribute"]] = r["attribute_value"]

	gallery_map: dict[str, list[str]] = {}
	for r in frappe.get_all(
		"Listing Image",
		filters={"parent": ["in", names], "parenttype": "Listing"},
		fields=["parent", "image", "sort_order"],
		order_by="sort_order asc",
	):
		if r.get("image"):
			gallery_map.setdefault(r["parent"], []).append(r["image"])

	variant_map: dict[str, list[dict]] = {}
	for r in frappe.get_all(
		"Listing Variant Item",
		filters={"parent": ["in", names], "parenttype": "Listing"},
		fields=[
			"parent",
			"variant_sku",
			"attribute_type",
			"attribute_value",
			"attribute_type_2",
			"attribute_value_2",
			"attribute_type_3",
			"attribute_value_3",
			"variant_price",
			"variant_stock",
			"idx",
		],
		order_by="idx asc",
	):
		variant_map.setdefault(r["parent"], []).append(r)

	columns = build_template_columns()
	headers = [c[0] for c in columns]
	canonicals = [c[1] for c in columns]

	rows: list[list[str]] = []
	for name in names:
		listing = by_name.get(name)
		if not listing:
			continue
		parent_map = _build_parent_value_map(
			listing, attr_map.get(name, {}), gallery_map.get(name, [])
		)
		rows.append(_row_from_map(canonicals, parent_map, listing=listing))
		# Varyant satırları (varsa) parent'tan sonra eklenir.
		for v in variant_map.get(name, []):
			vmap = _build_variant_value_map(listing.get("seller_sku") or "", v)
			rows.append(_row_from_map(canonicals, vmap, listing=None))

	if format == "xlsx":
		from openpyxl import Workbook
		from openpyxl.styles import Font, PatternFill

		wb = Workbook()
		ws = wb.active
		ws.title = "Products"
		ws.append(headers)
		for cell in ws[1]:
			cell.font = Font(bold=True)
			cell.fill = PatternFill("solid", fgColor="EDE9FE")
		for row in rows:
			ws.append(row)
		buf = io.BytesIO()
		wb.save(buf)
		content = buf.getvalue()
		file_name = "urunlerim_export.xlsx"
	else:
		buf = io.StringIO()
		writer = csv.writer(buf)
		writer.writerow(headers)
		writer.writerows(rows)
		content = b"\xef\xbb\xbf" + buf.getvalue().encode("utf-8")
		file_name = "urunlerim_export.csv"

	frappe.local.response.filename = file_name
	frappe.local.response.filecontent = content
	frappe.local.response.type = "binary"
